#!/usr/bin/env python3
"""Überführt die ABM-Projektliste (Excel) in eine SQL-Datei für Supabase.

Aufruf:
    python3 tools/import_excel.py "Projektliste_2026_digital.xlsx"
    -> tools/out/import.sql     (in Supabase "SQL Editor" einspielen)
    -> tools/out/bericht.md     (was korrigiert wurde, was geprüft werden sollte)

Die Dateien in tools/out/ enthalten echte Kundendaten und sind per .gitignore
ausgeschlossen. Niemals ins Repository committen.

Das Skript ist für den einmaligen Start gedacht. Die IDs werden aus den Inhalten
abgeleitet, ein zweiter Lauf legt daher keine Doppelten an (on conflict do nothing),
ändert aber auch nichts an bereits importierten Zeilen.
"""
from __future__ import annotations

import argparse
import collections
import datetime as dt
import re
import sys
import uuid
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

NAMESPACE = uuid.UUID("5d0f4c0e-6a39-4f3a-9d3e-0a6b1c2d3e4f")

MONTHS = {
    "januar": 1, "februar": 2, "märz": 3, "maerz": 3, "april": 4, "mai": 5, "juni": 6,
    "juli": 7, "august": 8, "september": 9, "oktober": 10, "november": 11, "dezember": 12,
}

# Spalten im Blatt "Blatt1" (1-basiert)
COL = dict(nr=1, firma=2, abm_nr=3, bauherr=4, ort=5, verkaeufer=6, architekt=7, planer=8,
           plangespraech=9, baugesuch=10, vertrag_versendet=11, vertrag_zurueck=12,
           genehmigung=13, baubetreuung=14)

STEP_COLUMNS = [
    ("plangespraech", "plangespraech"),
    ("vertrag_versendet", "vertrag_versendet"),
    ("vertrag_zurueck", "vertrag_zurueck"),
    ("baugesuch_uebergabe", "baugesuch"),
    ("baugenehmigung", "genehmigung"),
]

PLAUSIBLE_YEARS = range(2024, 2028)


# ---------------------------------------------------------------------------
# Bereinigung (rein, ohne Excel-Zugriff, daher gut testbar)
# ---------------------------------------------------------------------------
def clean_text(value) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def normalize_planer(value) -> tuple[str | None, bool]:
    """Gibt (Name, ist_storno_hinweis) zurück. 'L.Tas' wird zu 'L. Tas'."""
    text = clean_text(value)
    if text is None:
        return None, False
    if "storno" in text.casefold():
        return None, True
    return re.sub(r"\.(?=\S)", ". ", text), False


def normalize_baubetreuung(value) -> str | None:
    text = clean_text(value)
    if text is None:
        return None
    return text.rstrip(".").strip() or None


def parse_date(value) -> tuple[dt.date | None, str | None]:
    """Liefert (Datum, Hinweis). Hinweis ist gesetzt, wenn korrigiert werden musste."""
    if value is None or value == "":
        return None, None
    if isinstance(value, dt.datetime):
        return value.date(), None
    if isinstance(value, dt.date):
        return value, None
    if isinstance(value, (int, float)):
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=int(value))).date(), "Excel-Seriennummer umgerechnet"
    text = str(value).strip()
    m = re.fullmatch(r"(\d{1,2})\.(\d{1,2})\.(\d{1,4})", text)
    if m:
        day, month, year = int(m.group(1)), int(m.group(2)), m.group(3)
        note = None
        if len(year) < 4 or int(year) < 1000:
            # Tippfehler wie "206" oder "0206" statt "2026"
            if year.lstrip("0") in ("206", "26", "6", "2026"):
                year_i = 2026
            else:
                return None, f"Jahr nicht lesbar: {text!r}"
            note = f"Jahr korrigiert: {text!r}"
        else:
            year_i = int(year)
        try:
            return dt.date(year_i, month, day), note
        except ValueError:
            return None, f"Datum ungültig: {text!r}"
    return None, f"Datum nicht lesbar: {text!r}"


def parse_month_header(text: str) -> dt.date | None:
    m = re.fullmatch(r"\s*([A-Za-zäöüÄÖÜ]+)\s+(\d{4})\s*", text)
    if not m:
        return None
    month = MONTHS.get(m.group(1).casefold())
    return dt.date(int(m.group(2)), month, 1) if month else None


def detect_storno(row_texts: list[str], red_cells: int) -> bool:
    return red_cells >= 10 or any("storno" in t.casefold() for t in row_texts)


def uid(*parts: str) -> str:
    return str(uuid.uuid5(NAMESPACE, "|".join(p or "" for p in parts)))


def q(value) -> str:
    """SQL-Literal."""
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, dt.date):
        return f"'{value.isoformat()}'"
    return "'" + str(value).replace("'", "''") + "'"


# ---------------------------------------------------------------------------
# Datenstrukturen
# ---------------------------------------------------------------------------
@dataclass
class Project:
    source_row: int
    abm_nr: str | None
    kind: str
    status: str
    company: str | None
    salesperson: str | None
    bauherr_label: str
    location: str | None
    architect: str | None
    planer: str | None
    baubetreuung: str | None
    intake_month: dt.date | None
    steps: dict[str, dt.date] = field(default_factory=dict)
    notes: list[str] = field(default_factory=list)
    bauherr_key: str | None = None  # Schlüssel in contacts, falls zugeordnet


@dataclass
class Contact:
    row: int
    name: str
    address: str | None
    phone: str | None
    email: str | None
    building_address: str | None
    surveyor: str | None
    order_no: str | None
    basement_builder: str | None
    advisor: str | None
    architect: str | None


class Report:
    def __init__(self) -> None:
        self.corrections: list[str] = []
        self.review: list[str] = []

    def fixed(self, row: int, msg: str) -> None:
        self.corrections.append(f"Zeile {row}: {msg}")

    def check(self, row: int, msg: str) -> None:
        self.review.append(f"Zeile {row}: {msg}")


# ---------------------------------------------------------------------------
# Einlesen
# ---------------------------------------------------------------------------
def count_red(ws, row: int) -> int:
    n = 0
    for c in range(1, 15):
        f = ws.cell(row, c).fill
        if f and f.fill_type == "solid" and f.fgColor.type == "rgb" and f.fgColor.rgb == "FFFF0000":
            n += 1
    return n


def read_projects(ws, report: Report) -> list[Project]:
    projects: list[Project] = []
    month: dt.date | None = None
    for r in range(4, ws.max_row + 1):
        cells = {k: ws.cell(r, c).value for k, c in COL.items()}
        a = cells["nr"]
        others_empty = all(ws.cell(r, c).value is None for c in range(2, 15))
        if isinstance(a, str) and others_empty:
            parsed = parse_month_header(a)
            if parsed:
                month = parsed
            continue
        bauherr = clean_text(cells["bauherr"])
        abm_raw = clean_text(cells["abm_nr"])
        if not bauherr and not abm_raw:
            continue  # Platzhalterzeile nur mit laufender Nummer

        texts = [t for t in (clean_text(ws.cell(r, c).value) for c in range(1, 15)) if t]
        red = count_red(ws, r)
        status = "storno" if detect_storno(texts, red) else "aktiv"
        if status == "storno" and red >= 10 and not any("storno" in t.casefold() for t in texts):
            report.fixed(r, "nur rot markiert, ohne Text. Als Storno übernommen.")

        kind = "neubau"
        abm_nr = abm_raw
        company = clean_text(cells["firma"])
        if company and company.casefold() == "bauvoranfrage":
            # Eintrag steht versehentlich in der Spalte "Fertighausfirma"
            kind, company = "bauvoranfrage", None
            report.fixed(r, "'Bauvoranfrage' stand in der Firmenspalte. Als Bauvoranfrage ohne Firma übernommen.")
        if abm_raw and abm_raw.casefold() == "bauvoranfrage":
            kind, abm_nr = "bauvoranfrage", None
        elif abm_raw and re.fullmatch(r"\d{4}-", abm_raw):
            report.check(r, f"ABM-Nummer unvollständig ({abm_raw}). Nummer leer gelassen, Rest im Notizfeld.")
            abm_nr = None

        location = clean_text(cells["ort"])
        if location and "storno" in location.casefold():
            location = None
        planer, planer_storno = normalize_planer(cells["planer"])
        if planer_storno:
            report.fixed(r, "Storno-Vermerk in der Spalte Planer entfernt.")

        architect = clean_text(cells["architekt"])
        if architect:
            architect = architect.upper()
            if architect not in ("BE", "MC"):
                report.check(r, f"unbekanntes Architekten-Kürzel {architect!r}")

        project = Project(
            source_row=r,
            abm_nr=abm_nr,
            kind=kind,
            status=status,
            company=company,
            salesperson=clean_text(cells["verkaeufer"]),
            bauherr_label=bauherr or "(ohne Namen)",
            location=location,
            architect=architect,
            planer=planer,
            baubetreuung=normalize_baubetreuung(cells["baubetreuung"]),
            intake_month=month,
        )
        if abm_raw and abm_nr is None and kind == "neubau":
            project.notes.append(f"ABM-Nr. in der Excel unvollständig: {abm_raw}")
        if bauherr is None:
            report.check(r, "kein Bauherr eingetragen")

        for step_key, col_key in STEP_COLUMNS:
            value, note = parse_date(cells[col_key])
            if note:
                (report.check if value is None else report.fixed)(r, f"{step_key}: {note}")
            if value is None:
                continue
            if value.year not in PLAUSIBLE_YEARS:
                report.check(r, f"{step_key}: Datum {value.isoformat()} außerhalb des erwarteten Zeitraums")
            project.steps[step_key] = value

        # Reihenfolge-Plausibilität
        sent, back = project.steps.get("vertrag_versendet"), project.steps.get("vertrag_zurueck")
        if sent and back and back < sent:
            report.check(r, "Vertrag zurück liegt vor Vertrag versendet")
        projects.append(project)
    return projects


def read_contacts(ws) -> list[Contact]:
    contacts: list[Contact] = []
    architect: str | None = None
    for r in range(2, ws.max_row + 1):
        name = clean_text(ws.cell(r, 1).value)
        if not name:
            continue
        rest_empty = all(ws.cell(r, c).value is None for c in range(2, 10))
        if rest_empty:
            low = name.casefold()
            architect = "BE" if "willi" in low else "MC" if "mehmet" in low else architect
            continue

        def cell(c: int) -> str | None:
            v = ws.cell(r, c).value
            if v is None:
                return None
            lines = [re.sub(r"[ \t]+", " ", x).strip() for x in str(v).splitlines()]
            return "\n".join(x for x in lines if x) or None

        contacts.append(Contact(
            row=r, name=" ".join(name.split()), address=cell(2), phone=cell(3), email=cell(4),
            building_address=cell(5), surveyor=cell(6), order_no=cell(7),
            basement_builder=cell(8), advisor=cell(9), architect=architect,
        ))
    return contacts


def tokens(text: str) -> set[str]:
    return {t for t in re.split(r"[\s/,&\-\n]+", text.casefold()) if len(t) >= 3 and t != "und"}


def match_contacts(projects: list[Project], contacts: list[Contact]) -> dict[int, int]:
    """Ordnet Projekte (source_row) Kontakten (row) zu, nur bei eindeutigem Treffer.

    Bedingung: Nachname passt UND der erste Wortteil des Bauorts kommt in Anschrift oder
    Bau-Anschrift vor. Lieber keine Zuordnung als eine falsche mit fremden Kontaktdaten.
    """
    result: dict[int, int] = {}
    for p in projects:
        if not p.location or p.bauherr_label.startswith("("):
            continue
        place = re.split(r"[\s\-]+", p.location.casefold())[0]
        if len(place) < 4:
            continue
        label_tokens = tokens(p.bauherr_label)
        candidates = [
            c for c in contacts
            if label_tokens & tokens(c.name)
            and place in f"{c.address or ''} {c.building_address or ''}".casefold()
        ]
        if len(candidates) == 1:
            result[p.source_row] = candidates[0].row
    return result


# ---------------------------------------------------------------------------
# SQL ausgeben
# ---------------------------------------------------------------------------
def build_sql(projects: list[Project], contacts: list[Contact], matches: dict[int, int]) -> str:
    out: list[str] = ["-- Automatisch erzeugt von tools/import_excel.py. ENTHÄLT ECHTE KUNDENDATEN.",
                      "begin;", ""]

    companies = sorted({p.company for p in projects if p.company})
    prefixes: dict[str, collections.Counter] = collections.defaultdict(collections.Counter)
    for p in projects:
        if p.company and p.abm_nr and re.match(r"\d{4}-", p.abm_nr):
            prefixes[p.company][p.abm_nr[:4]] += 1
    for name in companies:
        prefix = prefixes[name].most_common(1)[0][0] if prefixes[name] else None
        out.append("insert into public.companies (id, name, number_prefix) values "
                   f"({q(uid('company', name))}, {q(name)}, {q(prefix)}) on conflict do nothing;")

    sales = sorted({(p.company, p.salesperson) for p in projects if p.company and p.salesperson})
    for company, person in sales:
        out.append("insert into public.salespeople (id, company_id, name) values "
                   f"({q(uid('sales', company, person))}, {q(uid('company', company))}, {q(person)}) "
                   "on conflict do nothing;")

    for code in sorted({p.architect for p in projects if p.architect in ("BE", "MC")}):
        out.append("insert into public.team_members (id, code, name, role) values "
                   f"({q(uid('team', code))}, {q(code)}, {q('Architekt ' + code)}, 'architekt') "
                   "on conflict do nothing;")
    for name in sorted({p.planer for p in projects if p.planer}):
        out.append("insert into public.team_members (id, name, role) values "
                   f"({q(uid('team', name))}, {q(name)}, 'planer') on conflict do nothing;")

    out.append("")
    for c in contacts:
        notes = f"Betreut von {c.architect}" if c.architect else None
        out.append(
            "insert into public.bauherren (id, name, address, phone, email, building_address, surveyor, "
            "builder_order_no, basement_builder, customer_advisor, notes) values ("
            f"{q(uid('bauherr', str(c.row), c.name))}, {q(c.name)}, {q(c.address)}, {q(c.phone)}, {q(c.email)}, "
            f"{q(c.building_address)}, {q(c.surveyor)}, {q(c.order_no)}, {q(c.basement_builder)}, "
            f"{q(c.advisor)}, {q(notes)}) on conflict do nothing;")

    out.append("")
    contacts_by_row = {c.row: c for c in contacts}
    for p in projects:
        pid = uid("project", str(p.source_row), p.abm_nr or "", p.bauherr_label)
        bauherr_id = None
        if p.source_row in matches:
            c = contacts_by_row[matches[p.source_row]]
            bauherr_id = uid("bauherr", str(c.row), c.name)
        notes = "; ".join(p.notes) or None
        out.append(
            "insert into public.projects (id, abm_nr, kind, status, company_id, salesperson_id, bauherr_id, "
            "bauherr_label, location, architect_id, planer_id, baubetreuung, intake_month, notes, source_row) values ("
            f"{q(pid)}, {q(p.abm_nr)}, {q(p.kind)}, {q(p.status)}, "
            f"{q(uid('company', p.company)) if p.company else 'null'}, "
            f"{q(uid('sales', p.company, p.salesperson)) if p.company and p.salesperson else 'null'}, "
            f"{q(bauherr_id)}, {q(p.bauherr_label)}, {q(p.location)}, "
            f"{q(uid('team', p.architect)) if p.architect in ('BE', 'MC') else 'null'}, "
            f"{q(uid('team', p.planer)) if p.planer else 'null'}, "
            f"{q(p.baubetreuung)}, {q(p.intake_month)}, {q(notes)}, {p.source_row}) on conflict do nothing;")
        for step_key, day in p.steps.items():
            out.append("insert into public.project_steps (project_id, step_key, occurred_on) values "
                       f"({q(pid)}, {q(step_key)}, {q(day)}) on conflict do nothing;")

    out += ["", "commit;", ""]
    return "\n".join(out)


def build_report(projects: list[Project], contacts: list[Contact], matches: dict[int, int],
                 report: Report, source: str) -> str:
    by_status = collections.Counter(p.status for p in projects)
    by_kind = collections.Counter(p.kind for p in projects)
    by_company = collections.Counter(p.company or "(ohne Firma)" for p in projects)
    nrs = collections.Counter(p.abm_nr for p in projects if p.abm_nr)
    dups = {k: v for k, v in nrs.items() if v > 1}
    rows_by_nr = {k: [p.source_row for p in projects if p.abm_nr == k] for k in dups}

    lines = [f"# Importbericht", "", f"Quelle: `{source}`  ", f"Erstellt: {dt.date.today().isoformat()}", "",
             "Dieser Bericht nennt Zeilennummern der Excel, keine Namen.", "", "## Zusammenfassung", "",
             f"- Projekte: **{len(projects)}** (aktiv {by_status['aktiv']}, Storno {by_status['storno']})",
             f"- davon Bauvoranfragen: {by_kind['bauvoranfrage']}",
             f"- Kontakte aus Blatt 2: {len(contacts)}, davon automatisch einem Projekt zugeordnet: {len(matches)}",
             f"- Korrekturen automatisch vorgenommen: {len(report.corrections)}",
             f"- Bitte manuell prüfen: {len(report.review) + len(dups)}", "", "Projekte je Firma:", ""]
    lines += [f"- {k}: {v}" for k, v in by_company.most_common()]
    lines += ["", "## Automatisch korrigiert", ""]
    lines += [f"- {x}" for x in report.corrections] or ["- keine"]
    lines += ["", "## Bitte prüfen", ""]
    lines += [f"- {x}" for x in report.review]
    lines += [f"- ABM-Nummer {k} kommt {v}× vor (Zeilen {', '.join(map(str, rows_by_nr[k]))})" for k, v in dups.items()]
    if not report.review and not dups:
        lines.append("- nichts")
    lines += ["", "## Hinweise", "",
              "- Rechnungen und Beträge stehen nicht in der Excel und werden daher nicht importiert.",
              "- Kontakte aus Blatt 2 werden nur zugeordnet, wenn Nachname und Bauort eindeutig passen.",
              "  Nicht zugeordnete Kontakte sind trotzdem importiert und können in der App verknüpft werden.",
              "  Eine sehr niedrige Trefferzahl heißt meist, dass Blatt 2 andere (z. B. frühere) Projekte betrifft.",
              "- Die Architekten-Kürzel BE und MC werden als `Architekt BE` / `Architekt MC` angelegt.",
              "  Vollständige Namen bitte in der Tabelle `team_members` ergänzen.", ""]
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", type=Path, help="Pfad zur Projektliste (.xlsx)")
    ap.add_argument("--out", type=Path, default=Path("tools/out"), help="Ausgabeordner (Standard: tools/out)")
    args = ap.parse_args(argv)

    wb = openpyxl.load_workbook(args.xlsx)
    if "Blatt1" not in wb.sheetnames:
        print(f"Blatt 'Blatt1' nicht gefunden. Vorhanden: {wb.sheetnames}", file=sys.stderr)
        return 1
    report = Report()
    projects = read_projects(wb["Blatt1"], report)
    contacts = read_contacts(wb["Kontaktdaten Bauherren"]) if "Kontaktdaten Bauherren" in wb.sheetnames else []
    matches = match_contacts(projects, contacts)

    args.out.mkdir(parents=True, exist_ok=True)
    (args.out / "import.sql").write_text(build_sql(projects, contacts, matches), encoding="utf-8")
    (args.out / "bericht.md").write_text(build_report(projects, contacts, matches, report, args.xlsx.name),
                                         encoding="utf-8")
    print(f"{len(projects)} Projekte, {len(contacts)} Kontakte ({len(matches)} zugeordnet), "
          f"{len(report.corrections)} Korrekturen, {len(report.review)} Prüfhinweise.")
    print(f"Geschrieben: {args.out / 'import.sql'} und {args.out / 'bericht.md'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
